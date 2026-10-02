from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_block_items(document: Document):
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def paragraph_text(paragraph: Paragraph) -> str:
    text = paragraph.text.replace("\t", " ⇥ ").strip()
    return " ".join(text.split())


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    path = Path(sys.argv[1])
    document = Document(path)

    paragraphs = []
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph_text(paragraph)
        style = paragraph.style.name if paragraph.style else ""
        xml = paragraph._p.xml
        if text or "TOC" in xml or "PAGEREF" in xml or "PAGE" in xml:
            field_codes = [
                node.text or ""
                for node in paragraph._p.xpath(".//w:instrText")
            ]
            paragraphs.append(
                {
                    "index": index,
                    "style": style,
                    "text": text,
                    "has_page_break": 'w:type="page"' in xml,
                    "has_toc": " TOC " in xml,
                    "has_pageref": "PAGEREF" in xml,
                    "has_field": "<w:fldChar" in xml or "<w:instrText" in xml,
                    "field_codes": field_codes,
                }
            )

    tables = []
    for index, table in enumerate(document.tables):
        rows = []
        for row in table.rows[:4]:
            rows.append([" ".join(cell.text.split())[:180] for cell in row.cells])
        tables.append(
            {
                "index": index,
                "rows": len(table.rows),
                "columns": len(table.columns),
                "sample": rows,
            }
        )

    body = []
    for index, item in enumerate(iter_block_items(document)):
        if isinstance(item, Paragraph):
            text = paragraph_text(item)
            style = item.style.name if item.style else ""
            if text:
                body.append({"body_index": index, "kind": "paragraph", "style": style, "text": text})
        else:
            sample = " | ".join(
                " / ".join(" ".join(cell.text.split()) for cell in row.cells)
                for row in item.rows[:2]
            )
            body.append(
                {
                    "body_index": index,
                    "kind": "table",
                    "rows": len(item.rows),
                    "columns": len(item.columns),
                    "sample": sample[:360],
                }
            )

    result = {
        "path": str(path),
        "sections": len(document.sections),
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "inline_shape_count": len(document.inline_shapes),
        "paragraphs": paragraphs,
        "tables": tables,
        "body": body,
    }
    if len(sys.argv) > 2:
        output_path = Path(sys.argv[2])
        output_path.write_text(
            json.dumps(result, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        body_lines = []
        for item in body:
            if item["kind"] == "paragraph":
                body_lines.append(
                    f'{item["body_index"]:04d}\tP\t{item["style"]}\t{item["text"]}'
                )
            else:
                body_lines.append(
                    f'{item["body_index"]:04d}\tT\t{item["rows"]}x{item["columns"]}\t{item["sample"]}'
                )
        output_path.with_suffix(".body.txt").write_text(
            "\n".join(body_lines) + "\n",
            encoding="utf-8",
        )
    else:
        json.dump(result, sys.stdout, ensure_ascii=False, indent=2)
        print()


if __name__ == "__main__":
    main()
